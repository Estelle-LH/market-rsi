"""Build the user-approved checkpoint protocol with the bundled docx runtime."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("MARKET_RSI_LEARNING_CHECKPOINT_UPGRADE_2026-10-05.docx")
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.75)
for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet", "List Number"):
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.08
doc.styles["Title"].font.size = Pt(21)
doc.styles["Title"].paragraph_format.space_after = Pt(10)
for name, size in (("Heading 1", 14), ("Heading 2", 12)):
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(14)
    doc.styles[name].paragraph_format.keep_with_next = True
doc.core_properties.title = "Market RSI Learning Checkpoint Upgrade"
doc.core_properties.subject = "Approved protocol for learning assessment and research branch continuation"
doc.core_properties.author = "Market RSI Supervisor"


def p(text, style=None):
    paragraph = doc.add_paragraph(text, style)
    paragraph.paragraph_format.widow_control = True
    return paragraph


def heading(text):
    doc.add_heading(text, 1)


def bullets(items):
    for text in items:
        p(text, "List Bullet")


def table(headers, rows, widths):
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    for col, width in zip(tab.columns, widths):
        col.width = Inches(width)
    for index, text in enumerate(headers):
        tab.rows[0].cells[index].text = text
    for row in rows:
        for cell, text in zip(tab.add_row().cells, row):
            cell.text = text
    repeat = OxmlElement("w:tblHeader")
    tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, row in enumerate(tab.rows):
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for col_index, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                edge = OxmlElement("w:" + side)
                for key, value in (("val", "single"), ("sz", "5"), ("color", "D9D9D9")):
                    edge.set(qn("w:" + key), value)
                borders.append(edge)
            props.append(borders)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                margin = OxmlElement("w:" + side)
                margin.set(qn("w:w"), "100")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.append(margins)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "243F60" if row_index == 0 else ("F3F6F9" if row_index % 2 else "FFFFFF"))
            props.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if row_index == 0 else "000000")
    p("").paragraph_format.space_after = Pt(0)


doc.add_paragraph("Market RSI Learning Checkpoint Upgrade", "Title")
p("Supervisor to Estelle Zhang   October 5 2026", "Subtitle")
p("We will assess each research child on three separate decisions: whether it replaces the best predictor, what evidence or capability it contributes, and whether it deserves another bounded experiment. A child does not need an immediate forecast-score improvement to remain useful. It needs a valid experiment, an honest evidence claim and a specific reason for its next allocation.")
p("Approved scope is local implementation and verification of this protocol. Historical scores, judgments, artifacts and exhausted budgets remain unchanged. The completed six-attempt and 24-fit batch stays closed. A live pilot requires a separately bounded next batch; this document does not grant new data, account calls, paid providers, release or promotion rights.")

heading("Current defect and decisions to separate")
p("The current recorder couples credit 0 to invalidity, credit 1 to inconclusive results, and credit 2 to support or refutation. That prevents accurate descriptions of valid experiments that add little knowledge, or useful findings whose performance evidence remains inconclusive. Valid REVERT candidates can already become parents; this upgrade makes their learning value and continuation rationale explicit.")
bullets([
    "Prediction decision: retain the frozen KEEP or REVERT judgment for replacing the incumbent.",
    "Learning decision: record new evidence, capability, uncertainty and demonstrated reuse separately.",
    "Exploration decision: continue, branch, repair, park or stop the exact recipe under a bounded allocation.",
])
p("No blended learning reward will be added to Brier or log loss. A useful finding cannot cancel poor predictions; a better prediction does not by itself demonstrate a better researcher. New rules apply prospectively and preserve old replay behavior.")

heading("Checkpoint evaluation axes")
table(
    ["Axis", "Question", "Evidence required"],
    [
        ("Validity", "Was the intended experiment executed correctly?", "Source and runtime, data-time checks, outputs, sample alignment and independent verification."),
        ("Prediction", "What changed versus parent and market?", "Frozen proper scores, paired chronological blocks, uncertainty and coverage."),
        ("Knowledge", "What did this distinguish or establish?", "A specific non-duplicate finding with its scope, limitations and unresolved alternatives."),
        ("Capability", "Did a tool or workflow solve a real problem?", "Demonstrated behavior, preserved correctness and an identified downstream use."),
        ("Reuse", "Did earlier evidence affect this step and help?", "Linked prior finding, changed action and observed consequence; citation alone is insufficient."),
        ("Efficiency and autonomy", "What did it cost and who intervened?", "End-to-end time, resources, failures and separate human and AI Supervisor involvement."),
    ],
    [1.2, 2.0, 3.8],
)
p("Execution failure is not scientific refutation. Preserve its record, mark performance as not evaluated, and assess any successful repair under a new version. Invalid or leaking results cannot supply valid performance evidence.")

heading("The checkpoint card")
p("Before execution, freeze the actual research parent, incumbent comparator and triggering evidence as distinct references. Identify the bounded research question, competing explanations, change layer and what observations would support, weaken or leave the question unresolved. Include the resource ceiling, allowed operations and protected identities.")
p("After execution, append the actual result and validity review, verified findings and remaining uncertainty, prediction decision, exploration decision, next discriminating experiment, authorship, interventions and actual cost. A checkpoint must not retrospectively invent the hypothesis it supposedly tested.")
p("Changes are classified as predictor C, harness H, researcher policy or memory R, with the protected evaluation kernel K and base model or runtime M separately identified. Composite Discovery changes are permitted, but their effect belongs to the combination until an ablation distinguishes contributors. Smallness means a reviewable causal and operational scope, not just a line count.")

heading("Evidence and demonstrated learning")
table(
    ["Label", "Requirement"],
    [
        ("Verified finding", "Independent evidence supports the finding within its stated scope."),
        ("Observed reuse", "A later decision actually changes because of that finding."),
        ("Demonstrated benefit", "Reuse improves a subsequent outcome, such as avoiding a known failure, shortening execution or enabling a valid experiment."),
        ("Transfer demonstrated", "A controlled comparison shows benefit on another period or task."),
    ],
    [1.75, 5.25],
)
p("A reference to an earlier lesson establishes neither benefit nor transfer. Evidence labels may remain pending or inconclusive. Changes to the researcher model weights are not required for learning through memory, selection or workflow, but any claimed benefit needs observation.")
p("Retain simple question-level research credit: 0 means no new verifiable evidence; 1 means a relevant non-duplicate finding; 2 means a valid test of a key hypothesis or a verified reusable capability. Supervisor verifies the artifacts and limitations. Credit is independent of validity, prediction replacement, support or refutation, and exploration eligibility. Repeated wording, duplicate evidence and cosmetic question renaming do not earn additional credit.")

heading("Branch allocation and renewal")
p("Keep the incumbent separate from a global pool of two or three active research branches and an archive of all candidates. At most two candidate training branches run concurrently. The active pool is selected globally, not three descendants per parent or a mechanical score ranking.")
p("A branch may receive another bounded allocation because of predictive promise, a verified useful capability, a distinct unresolved hypothesis or a reasonable low-cost first test. Exploration does not require statistical significance. Credit informs the decision alongside method diversity, remaining uncertainty and expected cost; it does not grant unlimited descendants.")
bullets([
    "Continue develops the current hypothesis; branch tests a materially different explanation while preserving the finding.",
    "Repair addresses implementation failure without treating it as performance evidence. The failed version remains preserved and its budget remains consumed.",
    "Park retains the source and knowledge but spends no current budget. Stop exact recipe ends unjustified repeats without erasing a whole information family.",
    "Each renewal states a concrete new question, expected evidence and cost. Stale branches return to the archive; reasonable untested hypotheses can receive a first small allocation.",
])
p("Approximately 30 percent exploration remains an adjustable project starting allocation. It is neither a scientific threshold nor a mandatory permanent quota. Historical consumed follow-up allowances, including C2, are not reset by a new rule.")

heading("D5 worked example")
p("The possession-only D5 experiment passed validity checks but received REVERT because both losses worsened versus C7. The verified finding is that positive stable fitted possession coefficients did not produce better check predictions for this recipe. Possession information generally remains unresolved; the result does not falsify the entire family.")
p("The reusable lesson is that coefficient stability alone is insufficient evidence for selecting a candidate. Preserve D5 and its predictions in the archive. A descendant deserves funding only if it tests a distinct justified question. Learning benefit remains pending until a later experiment uses this lesson with an observed consequence. These are development annotations; D5's original score and frozen labels are not rewritten.")

heading("Implementation sequence")
p("First implement opt-in assessment records and independent decisions in the existing recorder. Preserve old schemas, replay, scoring and permission checks. Next project verified findings, unresolved questions, reuse outcomes and complete costs into the existing Controller feedback path, and record why a branch receives its next allocation. Do not build a separate reward service or reinforcement-learning training system.")
p("Independent review checks the exact changed source, semantic scope and matched behavior before future activation. Over 200 changed production lines or more than two production modules triggers splitting or an explicit inseparability review. Passing below that threshold is not proof of safety. Keep the core recorder and feedback integration as separate bounded checkpoints.")

heading("Required verification")
bullets([
    "A valid REVERT parent can produce a child; a valid but uninformative result is not labelled invalid.",
    "An unresolved question can receive a bounded first test. Learning credit affects selection without modifying forecast scores or granting authority.",
    "Duplicate findings cannot accumulate credit; claimed reuse and benefit require linked evidence and independent verification.",
    "Failed execution cannot masquerade as negative performance evidence. A repair is a new version and does not erase consumed attempts.",
    "Global branch capacity, lineage, restart recovery and old-schema replay remain correct.",
    "Fixed scoring, data-time boundaries, protected Dev and Final, fees and budgets remain unchanged.",
])
p("Run targeted synthetic tests, the full relevant existing suite and an independent integrated review. Synthetic lifecycle tests validate the implementation, not prediction gains, unattended operation or the scientific superiority of self-evolution.")

heading("Live acceptance and research comparison")
p("After implementation acceptance, target three or four real prediction experiments within a separately approved next-batch budget. Include feedback-dependent continuation and a later reuse check. Success means valid execution and verifiable accumulation; it does not require every candidate to improve a score. Report all failures, actual costs and intervention sources.")
p("Then compare fixed and evolving research processes with the same base model, starting information, permissions, active-branch capacity and total budget. Both can use feedback, memory and multiple branches; only the evolving group changes the declared research process. Count implementation and review overhead, run failures and independent repetitions. Memory and exploration incentives need separate controls for attribution.")
p("Every retained child must state its contribution, evidence for another allocation and how that evidence affects subsequent action. Immediate statistical improvement is one possible contribution, not the admission requirement. Reused Train remains Discovery. Independent forecast confirmation requires separately frozen evaluation; protected data, external acquisition, paid calls, release and promotion remain outside this implementation scope.")

heading("Project choices and evidence limits")
p("The credit rubric, two-to-three branch capacity, exploration allocation and pilot size are project choices. Paired evaluation, information-availability checks and separation of adaptive Discovery from independent confirmation guide our assessment, but this document does not claim a new literature review or universal numerical thresholds. The upgrade is a user-approved Supervisor-directed harness change, not an already demonstrated autonomous researcher improvement.")
p("Implementation reference: learning-checkpoint-upgrade-20261005-v1, parent commit 83461bd. Supporting evidence is retained in the existing HUMAN_PROGRESS and independent D5 reviewer records. The old batch has five scientific experiments plus one parity run, all 24 fits completed, and raw market remains incumbent.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("Page ")
run.font.size = Pt(9)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
for border in doc.styles.element.xpath(".//w:pBdr") + doc.element.xpath(".//w:pBdr"):
    border.getparent().remove(border)
doc.save(OUTPUT)
print(OUTPUT)
