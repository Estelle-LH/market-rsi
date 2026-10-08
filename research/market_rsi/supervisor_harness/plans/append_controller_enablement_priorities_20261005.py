"""Append the approved priorities while retaining the historical Word content."""
from io import BytesIO
from pathlib import Path
import hashlib
import subprocess

from docx import Document
from docx.shared import Pt, RGBColor


REPO = Path(__file__).resolve().parents[4]
RELATIVE = "research/market_rsi/supervisor_harness/plans/MARKET_RSI_REPAIR_DISCOVERY_APPROVED_2026-10-05.docx"
PARENT = "938f6960471167b6dc0d21093c240873b3ccfbc4"
ORIGINAL_SHA256 = "87fde22cfcd2691f2c82a28e3a32074c02ba714d258656d337f773c48cb1bdb9"
TITLE = "Controller Enablement Priorities"
FRONT_NOTE = "   Historical record Budgets not renewed See current addendum"
PAGES = [
    [
        ("h1", TITLE),
        ("p", "Approved implementation addendum dated October 5 2026. Estelle instructed the Supervisor to add these priorities to this document and start. We will enable useful research tools and automatic experimental continuation before optional framework work. This addendum applies prospectively; the preceding approvals, execution windows and results remain historical records, not renewed budgets."),
        ("p", "The current verified source is 938f696. The latest continuation batch is closed at two candidates, eight fits and two original Controller decisions. Both candidates were independently checked and received REVERT; raw market remains incumbent. The trajectory demonstrates feedback-dependent candidate selection, but substantial AI Supervisor coordination remains. It does not demonstrate unattended continuation, autonomous research-process modification or process superiority."),
        ("h2", "P0 Enable scoped scientific capability"),
        ("p", "Replace the decision-only tool prohibition with scoped research tools. The Controller may read approved code, history, predictions and memory, investigate resident Train, create candidate patches in separate research directories, request reviewed analysis and local training, retrieve results and propose versioned research-tool or workflow changes. Methods remain open, including combined predictor changes. Tool permissions constrain access and resource use, not the algorithm menu."),
        ("p", "Keep protected Dev and Final, the scorer, frozen samples and labels, credentials, budget authority and historical artifacts outside Controller modification. Local Train access does not automatically authorize transferring raw data to an account or provider. Enforce approved destinations and payload scope before transfer. Broker tools must reject forbidden access before execution; an instruction or post-event rejection is insufficient. A read-only filesystem sandbox alone does not prevent protected reads."),
        ("p", "Start with a hash-bound evidence-reading bridge, then integrate it with the account consumer in a separate checkpoint. The installed CLI reports version 0.160.0 and supports local configuration overrides. MCP allowlisting alone does not prove that shell, browser or other built-in channels are absent. Demonstrate the effective tool catalog and access restrictions before a live tool session."),
        ("h2", "P0 Separate safety from research allocation"),
        ("p", "Execution eligibility depends on valid source and artifacts, permitted operations and remaining resources. Research priority depends on evidence, method differences, unresolved questions and cost. Incumbent replacement follows the unchanged prediction rule. Research credit informs allocation but is not a universal prerequisite for a reasonable first hypothesis."),
        ("p", "Source inspection confirms that v4 already permits a valid zero-credit REVERT parent one bounded follow-up. Do not reimplement that capability or claim all REVERT parents are forbidden. Review the remaining one-follow-up and zero-or-duplicate-credit continuation restrictions prospectively. A distinct justified question may deserve another bounded allocation; exhausted old allowances and attempts stay consumed. Invalid results cannot provide performance evidence."),
        ("h2", "P0 Remove unnecessary legacy dependencies"),
        ("p", "E2B remains retired. Tinker is an optional legacy adapter disabled by default, not a mainline credential, SDK or admission requirement. Verify the existing account and local-training route works without either service. Neither ran in the latest batch, so dependency cleanup must not replace the actual tool and handoff repairs. Preserve legacy code and receipts; do not delete external resources or activate paid fallbacks."),
    ],
    [
        ("h1", "P1 Complete unattended real experiments"),
        ("p", "Connect the existing preparation, launch and acceptance operations into one thin continuation driver. The sequence is verified evidence, original Controller decision, candidate assignment, implementation, independent review, training, independently verified feedback and the next Controller decision. Mechanical fields such as IDs, paths, source bindings and reservations come from the driver. Scientific hypotheses and continuation choices come from the Controller."),
        ("p", "Generalize the existing evidence adapter beyond its C1 and C7 parent formats. Verified saved predictions provide a model-independent comparison interface; inheriting fitted model state is optional. Preserve identical task, rows, labels, folds, information times and scoring. Record the actual research parent separately from the incumbent and market baseline."),
        ("p", "Return failures as factual observations without fabricated predictions or scores. Failed experiments consume their allocated resources. The Controller decides whether a bounded repair, another branch or stopping the exact recipe is warranted. A repair creates a new version and preserves the failed record. Never automatically resample an uncertain original account call."),
        ("p", "P1 acceptance requires two consecutive real candidates with keyed predictions and independently verified scores, without a human continue message or hand-edited operational script between rounds. The second original Controller choice must identify evidence from the first experiment that changed its decision. Two preselected independent models do not satisfy this check. A negative score is not a batch stop condition."),
        ("h1", "P2 Sustain useful exploration and recovery"),
        ("p", "Maintain the incumbent, a global active pool of two or three justified branches, and the full archive separately. Select the pool globally using evidence, method diversity, specific unresolved hypotheses and expected cost. Do not mechanically keep the score-top-three or three children per parent. Preserve valid losing candidates and their lessons, with explicit reasons to retain, park, branch, repair or stop the exact recipe."),
        ("p", "Research credit must visibly affect later choices or allocations without changing prediction scoring or granting authority. Credit zero denotes no new checkable evidence; one denotes a relevant nonduplicate finding; two denotes a valid key-hypothesis test or demonstrated reusable capability. Verify the underlying evidence independently. A useful counterexample can justify redirecting effort rather than repeating a disproven recipe."),
        ("p", "Extend the existing worker and journal for bounded recovery. Reconcile completed work after interruption, identify live jobs reliably, prevent duplicate fits or account calls, and retain partial failures. Test interruptions before launch, during execution, after output creation and before feedback acknowledgement. Recovery must preserve both artifacts and consumed resources. Broker restrictions are not a claim of arbitrary-code OS containment."),
        ("p", "At most two candidate branches run concurrently with controlled threads, shared read-only data and separate code and artifact directories. Pin each branch source and comparator. Root alone integrates changes and writes shared budget, pool, incumbent and experiment state. Return independently verified feedback asynchronously; do not wait for unrelated branches or invent unfinished results."),
    ],
    [
        ("h1", "P3 Improve memory and checkpoint assessment"),
        ("p", "Replace recursive full-history payloads with a compact current packet and retrieval of immutable original evidence. The latest two account decisions consumed 333112 input tokens, so context construction is an operational cost worth measuring. Retain findings, scope, conditions, uncertainty, negative results and their consequences. Measure retrieval correctness, input tokens, repeated mistakes and time to a valid experiment; shorter context alone is not stronger research."),
        ("p", "Return Brier, log loss, calibration with counts, paired parent and market differences, date and week results, grouped uncertainty and relevant residual diagnostics. Binary-outcome Brier equals MSE under identical weighting; label this explicitly instead of presenting a duplicate metric as independent evidence. Future-price MSE belongs to a separate task. Do not modify old scores, KEEP thresholds or sample selection after observing results."),
        ("p", "Assess execution validity, scientific learning, continued exploration value, prediction performance and operational benefit separately. A useful negative finding need not improve predictions immediately. Credit cannot cancel prediction loss. Every retained branch needs a concrete question and proportionate next allocation; more code, calls or logs do not establish learning."),
        ("p", "Allow Controller-authored, versioned proposals for research tools, memory and selection procedures. Review, test and activate accepted versions at an idle checkpoint. A previous lesson affecting a later choice demonstrates observed reuse; its benefit still needs evidence. Do not force researcher-code modifications simply to register an evolution event."),
        ("h1", "P4 Test research mechanisms and independent gains"),
        ("p", "After continuation is reliable, compare fixed and evolving research processes with the same base-model version, starting information, tools, permissions, active-pool capacity and total budget. Both arms can use ordinary memory, feedback and repeated experimentation; only the declared process-modification capability differs. Include implementation, review, failures and interventions in resources. A single trajectory is a pilot, not superiority evidence. Test exploration incentives separately to avoid confounded attribution."),
        ("p", "Repeatedly viewed Train remains Discovery. Freeze selection and submission before new evaluation results. Independently confirm forecasts on unselected events, preferably events that occur and settle after candidate and model-version freeze. Historical event time does not establish real-time availability. Protected evaluation, new data acquisition and live literature retrieval remain separately authorized capabilities."),
    ],
    [
        ("h1", "Execution checkpoints and resource authority"),
        ("p", "Immediate order is P0 then P1. Later memory, recovery and mechanism work must not postpone the first real pair. Each implementation checkpoint changes one coherent component, records its actual author and C predictor, H harness or R researcher-policy layer, and retains exact source, tests, independent review and rollback. Approximately 200 changed lines or more than two production modules is a split or inseparability-review trigger, not a scientific threshold or safety guarantee. Composite Discovery is allowed with appropriately limited attribution."),
        ("p", "Use existing Git checkpoints and records, not a new reward service, benchmark, dashboard or approval system. Parallelize only non-overlapping ready engineering work and independent review. Document the actual parent, triggering evidence, proposer, change, execution, result, decision and subsequent use. Human-directed Supervisor repairs are harness work, not autonomous evolution."),
        ("p", "Start code and synthetic verification under the current instruction. The first evidence-tools checkpoint has zero actual Train reads, fits, live Controller transport calls, provider calls and network requests. All prior batch limits remain exhausted. A suggested later acceptance batch of 90 minutes, three candidate attempts and twelve fits is a proposed resource envelope, not an authorization. Record fresh batch-level account and payload authority before launch, with no per-round method approval inside that scope. Count failed attempts and all account and helper usage; unavailable total monetary cost remains unknown."),
        ("p", "These capacities, quotas and acceptance sizes are project choices. The addition does not report a new literature review, confirmed market edge, research-process superiority or completed live enablement. Those claims require their respective execution and comparison evidence."),
    ],
]


def append():
    output = REPO / RELATIVE
    original = subprocess.check_output(["git", "show", f"{PARENT}:{RELATIVE}"], cwd=REPO)
    if hashlib.sha256(original).hexdigest() != ORIGINAL_SHA256:
        raise ValueError("original approved document commitment changed")
    current = Document(output)
    if hashlib.sha256(output.read_bytes()).hexdigest() != ORIGINAL_SHA256 and TITLE not in [p.text for p in current.paragraphs]:
        raise ValueError("unrelated current document edit; do not overwrite")
    doc = Document(BytesIO(original))
    original_paragraphs = [p.text for p in doc.paragraphs]
    original_tables = [[list(c.text for c in row.cells) for row in table.rows] for table in doc.tables]
    note = doc.paragraphs[1].add_run(FRONT_NOTE)
    note.font.name = "Arial"
    note.font.size = Pt(10.5)
    note.font.color.rgb = RGBColor(0, 0, 0)
    for page in PAGES:
        doc.add_page_break()
        for kind, text in page:
            paragraph = doc.add_heading(text, 1 if kind == "h1" else 2) if kind != "p" else doc.add_paragraph(text)
            paragraph.paragraph_format.widow_control = True
            if kind == "p":
                paragraph.paragraph_format.space_after = Pt(7)
                paragraph.paragraph_format.line_spacing = 1.08
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.size = Pt(15 if kind == "h1" else 12 if kind == "h2" else 11)
    doc.core_properties.comments = "User-approved Controller enablement priorities appended October 5 2026; historical approvals and budgets not renewed."
    doc.save(output)
    checked = Document(output)
    retained = [p.text.removesuffix(FRONT_NOTE) for p in checked.paragraphs]
    assert retained[:len(original_paragraphs)] == original_paragraphs
    assert [[list(c.text for c in row.cells) for row in table.rows] for table in checked.tables] == original_tables
    print(f"Updated {output}; retained {len(original_paragraphs)} historical paragraphs and {len(original_tables)} tables")


if __name__ == "__main__":
    append()
